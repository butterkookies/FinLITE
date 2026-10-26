import os
import sys
import json
import copy
from datetime import datetime
import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_TEMPLATE = os.path.join(BASE_DIR, 'templates', 'LITE-Financial-Report-Universal-Template.docx')
DEFAULT_OUTPUT = os.path.join(BASE_DIR, 'exports', 'FinLITE-Exported-Report.docx')

def get_current_academic_calendar():
    """Dynamically determine current academic year and semester based on current date."""
    now = datetime.now()
    transmittal_date = now.strftime('%B %d, %Y')
    start_year = now.year if now.month >= 8 else now.year - 1
    academic_year = f"{start_year}–{start_year + 1}"
    semester = "1st SEMESTER" if (8 <= now.month <= 12 or now.month == 1) else "2nd SEMESTER"
    return transmittal_date, semester, academic_year

def ensure_invisible_table_borders(tbl):
    """Ensure all borders on the table are nil (invisible) so no box gridlines appear."""
    tblPr = tbl._element.tblPr
    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is None:
        tblBorders = parse_xml(r'<w:tblBorders %s>'
                               r'<w:top w:val="nil"/>'
                               r'<w:left w:val="nil"/>'
                               r'<w:bottom w:val="nil"/>'
                               r'<w:right w:val="nil"/>'
                               r'<w:insideH w:val="nil"/>'
                               r'<w:insideV w:val="nil"/>'
                               r'</w:tblBorders>' % nsdecls('w'))
        tblPr.append(tblBorders)
    else:
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = tblBorders.find(qn(f'w:{side}'))
            if b is not None:
                b.set(qn('w:val'), 'nil')
                if qn('w:sz') in b.attrib:
                    del b.attrib[qn('w:sz')]

def set_cell_border(cell, **kwargs):
    """
    Set individual cell borders: top, bottom, left, right.
    Example: set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": "000000"})
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b = OxmlElement(f'w:{edge}')
            for key, val in edge_data.items():
                b.set(qn(f'w:{key}'), str(val))
            tcBorders.append(b)

def replace_text_in_paragraph(p, search, replace):
    if search not in p.text:
        return
    full_text = p.text
    new_text = full_text.replace(search, replace)
    if p.runs:
        font_name = p.runs[0].font.name or 'Calibri'
        font_size = p.runs[0].font.size or Pt(11)
        bold = p.runs[0].bold
        italic = p.runs[0].italic
        p.text = new_text
        for r in p.runs:
            r.font.name = font_name
            r.font.size = font_size
            r.bold = bold
            r.italic = italic
    else:
        p.text = new_text

def replace_placeholders(doc, config):
    default_td, default_sem, default_ay = get_current_academic_calendar()
    sem = config.get('{{SEMESTER}}') or default_sem
    ay = config.get('{{ACADEMIC_YEAR}}') or default_ay
    for p in doc.paragraphs:
        if "This is to respectfully submit the attached" in p.text:
            p.text = f'This is to respectfully submit the attached “LITE Financial Report” for the {sem} Academic Year {ay}.'
            if p.runs:
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(11)
            continue

        for tag, val in config.items():
            if tag in p.text:
                replace_text_in_paragraph(p, tag, val)
                
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for tag, val in config.items():
                        if tag in p.text:
                            replace_text_in_paragraph(p, tag, val)

def resolve_typography(data):
    """Resolve font family and base size (Calibri 11pt default vs Times New Roman 12pt manuscript standard)."""
    if not isinstance(data, dict):
        return 'Calibri', 11
    typography = str(data.get('typography') or '').strip().lower()
    font_family = str(data.get('font_family') or '').strip()
    font_size_pt = data.get('font_size_pt')
    if typography == 'times' or font_family.lower() == 'times new roman':
        return 'Times New Roman', float(font_size_pt or 12)
    return (font_family or 'Calibri'), float(font_size_pt or 11)

def setup_footer(section, style="page_x_of_y", font_name="Calibri"):
    """
    Configure dynamic Word OpenXML footer fields.
    Options:
    - 'page_x_of_y': Page X of Y on all pages
    - 'page_x_of_y_no_cover': Page X of Y on pages 2-4, suppressed on cover/Page 1
    - 'simple': 1, 2, 3...
    - 'none': footer cleared
    """
    if style == 'none':
        for p in section.footer.paragraphs:
            p.text = ''
        return

    def populate_p(p, is_simple=False):
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        def make_run_xml(text):
            return parse_xml(r'<w:r %s><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="22"/><w:color w:val="555555"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (nsdecls('w'), font_name, font_name, text))
        def make_field_xml(field_type, sample_text):
            return parse_xml(r'<w:fldSimple %s w:instr="%s"><w:r><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="22"/><w:color w:val="555555"/></w:rPr><w:t>%s</w:t></w:r></w:fldSimple>' % (nsdecls('w'), field_type, font_name, font_name, sample_text))
        
        p.text = ''
        if is_simple:
            p._p.append(make_field_xml("PAGE", "1"))
        else:
            p._p.append(make_run_xml("Page "))
            p._p.append(make_field_xml("PAGE", "1"))
            p._p.append(make_run_xml(" of "))
            p._p.append(make_field_xml("NUMPAGES", "4"))

    if style == 'page_x_of_y_no_cover':
        section.different_first_page_header_footer = True
        for child in list(section.header._element):
            section.first_page_header._element.append(copy.deepcopy(child))
        for p in section.first_page_footer.paragraphs:
            p.text = ''
        p_main = section.footer.paragraphs[0]
        populate_p(p_main, is_simple=False)
    elif style == 'simple':
        section.different_first_page_header_footer = False
        p_main = section.footer.paragraphs[0]
        populate_p(p_main, is_simple=True)
    else:  # 'page_x_of_y'
        section.different_first_page_header_footer = False
        p_main = section.footer.paragraphs[0]
        populate_p(p_main, is_simple=False)

def create_borderless_signatory_table(doc, signatories_list, num_cols, font_name='Calibri', font_size_pt=11):
    """
    Build structured multi-column borderless table for student officers.
    Guarantees exact column widths, zero tab wrapping misalignment, and consistent typography.
    Guards against division by zero if num_cols <= 0.
    """
    if num_cols <= 0 or not signatories_list:
        return None

    table = doc.add_table(rows=1, cols=num_cols)
    ensure_invisible_table_borders(table)
    col_width = Inches(6.5 / num_cols)
    for col in table.columns:
        col.width = col_width
    row = table.rows[0]
    for idx, sig in enumerate(signatories_list):
        if idx >= num_cols:
            break
        cell = row.cells[idx]
        cell.width = col_width
        
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
        tcPr.append(tcMar)
        
        # Consistent space for signature above officer name
        p_name = cell.paragraphs[0]
        p_name.paragraph_format.space_before = Pt(24)
        p_name.paragraph_format.space_after = Pt(1)
        p_name.paragraph_format.line_spacing = 1.0

        if sig.get('name'):
            r_name = p_name.add_run(sig['name'])
            r_name.font.name = font_name
            r_name.font.size = Pt(font_size_pt)
            r_name.bold = True
            
        p_role = cell.add_paragraph()
        p_role.paragraph_format.space_before = Pt(0)
        p_role.paragraph_format.space_after = Pt(4)
        p_role.paragraph_format.line_spacing = 1.0
        
        if sig.get('role'):
            r_role = p_role.add_run(sig['role'])
            r_role.font.name = font_name
            r_role.font.size = Pt(font_size_pt)

    # Fill remaining columns with empty cells to keep left-alignment without stretching
    for empty_idx in range(len(signatories_list), num_cols):
        empty_cell = row.cells[empty_idx]
        empty_cell.width = col_width
            
    tbl_elm = table._element
    tbl_elm.getparent().remove(tbl_elm)
    return tbl_elm

def create_borderless_noted_table(doc, advisers, admins, font_name='Calibri', font_size_pt=11):
    """
    Build unified 2-column borderless table for institutional oversight:
    Row 0: Club Advisers
    Row 1: Academic Administration (Program Director, Dean)
    """
    has_advisers = bool(advisers)
    has_admins = bool(admins)
    if not has_advisers and not has_admins:
        return None

    num_rows = (1 if has_advisers else 0) + (1 if has_admins else 0)
    table = doc.add_table(rows=num_rows, cols=2)
    ensure_invisible_table_borders(table)
    col_width = Inches(3.25)
    for col in table.columns:
        col.width = col_width

    curr_row_idx = 0

    if has_advisers:
        row = table.rows[curr_row_idx]
        for col_idx in range(2):
            cell = row.cells[col_idx]
            cell.width = col_width
            tcPr = cell._element.get_or_add_tcPr()
            tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
            tcPr.append(tcMar)
            if col_idx < len(advisers):
                sig = advisers[col_idx]
                p_name = cell.paragraphs[0]
                p_name.paragraph_format.space_before = Pt(24)
                p_name.paragraph_format.space_after = Pt(1)
                p_name.paragraph_format.line_spacing = 1.0
                r_name = p_name.add_run(sig.get('name', ''))
                r_name.font.name = font_name
                r_name.font.size = Pt(font_size_pt)
                r_name.bold = True

                p_role = cell.add_paragraph()
                p_role.paragraph_format.space_before = Pt(0)
                p_role.paragraph_format.space_after = Pt(4)
                p_role.paragraph_format.line_spacing = 1.0
                r_role = p_role.add_run(sig.get('role', ''))
                r_role.font.name = font_name
                r_role.font.size = Pt(font_size_pt)
        curr_row_idx += 1

    if has_admins:
        row = table.rows[curr_row_idx]
        for col_idx in range(2):
            cell = row.cells[col_idx]
            cell.width = col_width
            tcPr = cell._element.get_or_add_tcPr()
            tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
            tcPr.append(tcMar)
            if col_idx < len(admins):
                sig = admins[col_idx]
                p_name = cell.paragraphs[0]
                p_name.paragraph_format.space_before = Pt(20 if has_advisers else 24)
                p_name.paragraph_format.space_after = Pt(1)
                p_name.paragraph_format.line_spacing = 1.0
                r_name = p_name.add_run(sig.get('name', ''))
                r_name.font.name = font_name
                r_name.font.size = Pt(font_size_pt)
                r_name.bold = True

                p_role = cell.add_paragraph()
                p_role.paragraph_format.space_before = Pt(0)
                p_role.paragraph_format.space_after = Pt(4)
                p_role.paragraph_format.line_spacing = 1.0
                r_role = p_role.add_run(sig.get('role', ''))
                r_role.font.name = font_name
                r_role.font.size = Pt(font_size_pt)

    tbl_elm = table._element
    tbl_elm.getparent().remove(tbl_elm)
    return tbl_elm

def add_receipt_sheets_table(doc, receipts, academic_year, semester, font_name='Calibri', font_size_pt=11):
    """
    Renders 2-column Word table holding up to 4 receipts per page (2 rows x 2 cols).
    Scales pictures proportionally into max 3.1" x 3.0" bounding box.
    First receipt sheet displays 'Receipts' heading; subsequent sheets show only the 2x2 grid.
    """
    if not receipts:
        return

    from PIL import Image

    chunks = [receipts[i:i + 4] for i in range(0, len(receipts), 4)]

    for sheet_idx, chunk in enumerate(chunks):
        if sheet_idx > 0 or len(doc.paragraphs) > 0:
            p_break = parse_xml(r'<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls('w'))
            doc._element.body.append(p_break)

        if sheet_idx == 0:
            p_title = doc.add_paragraph()
            p_title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_before = Pt(4)
            p_title.paragraph_format.space_after = Pt(6)

            r_heading = p_title.add_run("Receipts")
            r_heading.font.name = font_name
            r_heading.font.size = Pt(font_size_pt)
            r_heading.bold = True

        table = doc.add_table(rows=2, cols=2)
        tblPr = table._element.tblPr
        tblBorders = parse_xml(r'<w:tblBorders %s>'
                               r'<w:top w:val="none"/>'
                               r'<w:left w:val="none"/>'
                               r'<w:bottom w:val="none"/>'
                               r'<w:right w:val="none"/>'
                               r'<w:insideH w:val="none"/>'
                               r'<w:insideV w:val="none"/>'
                               r'</w:tblBorders>' % nsdecls('w'))
        tblPr.append(tblBorders)

        for col in table.columns:
            col.width = Inches(3.2)

        for idx, rc in enumerate(chunk):
            r_idx = idx // 2
            c_idx = idx % 2
            cell = table.rows[r_idx].cells[c_idx]
            cell.width = Inches(3.2)

            tcPr = cell._element.get_or_add_tcPr()
            tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="80" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tcMar>' % nsdecls('w'))
            tcPr.append(tcMar)

            set_cell_border(cell,
                top={"val": "single", "sz": "4", "color": "D1D5DB"},
                bottom={"val": "single", "sz": "4", "color": "D1D5DB"},
                left={"val": "single", "sz": "4", "color": "D1D5DB"},
                right={"val": "single", "sz": "4", "color": "D1D5DB"}
            )

            p_img = cell.paragraphs[0]
            p_img.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)

            img_path = rc.get('local_path')
            if img_path and os.path.exists(img_path):
                try:
                    with Image.open(img_path) as im:
                        im_w, im_h = im.size
                    if im_w > 0 and im_h > 0:
                        aspect = im_w / im_h
                        # Max box: 3.1" wide by 3.0" tall
                        if aspect > (3.1 / 3.0):
                            w = Inches(3.1)
                            h = Inches(3.1 / aspect)
                        else:
                            h = Inches(3.0)
                            w = Inches(3.0 * aspect)
                        p_img.add_run().add_picture(img_path, width=w, height=h)
                except Exception as img_err:
                    p_img.add_run(f"[Receipt Image: {rc.get('title', 'Voucher')}]")
            else:
                p_img.add_run(f"[Receipt: {rc.get('title', 'Voucher')}]")

            p_cap = cell.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(2)
            p_cap.paragraph_format.line_spacing = 1.0

            r_d = p_cap.add_run(f"{rc.get('date', '')}  •  ")
            r_d.font.name = font_name
            r_d.font.size = Pt(9.5)

            r_a = p_cap.add_run(f"P{float(rc.get('amount', 0)):,.2f}\n")
            r_a.font.name = font_name
            r_a.font.size = Pt(9.5)
            r_a.bold = True

            r_t = p_cap.add_run(f"{rc.get('title', 'Disbursement Particulars')}\n")
            r_t.font.name = font_name
            r_t.font.size = Pt(9.5)
            r_t.bold = True

            r_c = p_cap.add_run(f"{rc.get('category', 'Disbursement')}")
            r_c.font.name = font_name
            r_c.font.size = Pt(8.5)
            r_c.font.color.rgb = docx.shared.RGBColor(100, 100, 100)

def enforce_calibri_11(doc, font_name='Calibri', font_size_pt=11):
    """Strictly enforce the selected typography (Calibri 11pt or Times New Roman 12pt) across all runs in paragraphs and table cells."""
    for style in doc.styles:
        if hasattr(style, 'font') and style.font is not None:
            style.font.name = font_name
            style.font.size = Pt(font_size_pt)

    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size_pt)

    for tbl in doc.tables:
        ensure_invisible_table_borders(tbl)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = font_name
                        r.font.size = Pt(font_size_pt)

    for s in doc.sections:
        for p in s.header.paragraphs:
            for r in p.runs:
                r.font.name = font_name
                r.font.size = Pt(font_size_pt)

def populate_report(data, template_path=DEFAULT_TEMPLATE, output_path=DEFAULT_OUTPUT):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Universal template not found at {template_path}")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = docx.Document(template_path)
    
    # 1. STRICT METADATA CONFORMANCE
    core = doc.core_properties
    core.last_modified_by = 'Andrei John Geronimo'
    core.revision = 3

    # Section Margins & Page Size (Letter: 8.5 x 11 in)
    # Calibrated Minimal Header Space & Top Margin to eliminate vertical misalignment
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.header_distance = Inches(0.20)
        s.top_margin = Inches(1.10)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.footer_distance = Inches(0.35)

    # Dynamic Bottom-Right Footer Page Numbers & Typography Resolution
    font_name, font_size_pt = resolve_typography(data)
    page_number_style = data.get('page_number_style', 'page_x_of_y')
    setup_footer(doc.sections[0], style=page_number_style, font_name=font_name)

    include_cc = bool(data.get('include_cc', False))
    cc_recipients = data.get('cc_recipients', ['cc: Co-Curricular', "cc: Dean's Office"])

    sig = data.get('signatories', {}) or {}
    summary = data.get('summary', {}) or {}

    def get_val(*keys, default=''):
        for k in keys:
            if k in data and data[k] is not None and str(data[k]).strip() != '':
                return str(data[k]).strip()
        return default

    # Transmittal & Meta
    default_transmittal_date, default_sem, default_ay = get_current_academic_calendar()
    transmittal_date = get_val('transmittal_date', 'transmittalDate', default=default_transmittal_date)
    coordinator_name = get_val('coordinator_name', 'coordinatorName', default='Ms. Ligaya H. Estrella')
    coordinator_role = get_val('coordinator_role', 'coordinatorRole', default='Co-curricular Affairs Coordinator')
    coordinator_salutation = get_val('coordinator_salutation', 'coordinatorSalutation', default='Mrs. Estrella')
    semester = get_val('semester', default=default_sem)
    academic_year = get_val('academic_year', 'academicYear', default=default_ay)
    period_desc = get_val('period_desc', 'periodDesc', 'eventName', 'activityTitle', default='Financial Operations')
    final_as_of = get_val('final_as_of', 'finalAsOf', default=f"As of {transmittal_date}")

    # Signatories Extraction (AY 2026-2027 & AY 2025-2026 Compatible)
    pres_name = get_val('president_name', default=sig.get('presidentName', sig.get('notedBy', 'GERONIMO, ANDREI JOHN P.'))).upper()
    pres_role = get_val('president_role', default=sig.get('presidentRole', sig.get('notedRole', 'LITE PRESIDENT'))).upper()
    vp_name = get_val('vice_president_name', 'vp_name', default=sig.get('vicePresidentName', sig.get('vpName', 'BAJAMUNDE, JHELLIAM T.'))).upper()
    vp_role = get_val('vice_president_role', 'vp_role', default=sig.get('vicePresidentRole', sig.get('vpRole', 'LITE VICE PRESIDENT'))).upper()
    tres_name = get_val('treasurer_name', default=sig.get('treasurerName', sig.get('preparedBy', 'CABANAG, ROSE MARIELLE P.'))).upper()
    tres_role = get_val('treasurer_role', default=sig.get('treasurerRole', sig.get('preparedRole', 'LITE TREASURER'))).upper()
    aud_name = get_val('auditor_name', default=sig.get('auditorName', sig.get('reviewedBy', 'SARA, RISHAME P.'))).upper()
    aud_role = get_val('auditor_role', default=sig.get('auditorRole', sig.get('reviewedRole', 'LITE AUDITOR'))).upper()
    
    adv1_name = get_val('adviser1_name', 'adviser1', 'adviser_name', default=sig.get('adviser1Name', sig.get('adviser1', sig.get('adviserName', 'MS. KIMBERLY DAWN JATULAN')))).upper()
    adv1_role = get_val('adviser1_role', default=sig.get('adviser1Role', sig.get('adviserRole', 'LITE CLUB ADVISER'))).upper()
    adv2_name = get_val('adviser2_name', 'adviser2', default=sig.get('adviser2Name', sig.get('adviser2', 'MS. KRIZIA MAE GENOVIA'))).upper()
    adv2_role = get_val('adviser2_role', default=sig.get('adviser2Role', 'LITE CLUB ADVISER')).upper()

    dir_name = get_val('director_name', default=sig.get('directorName', '')).upper()
    dir_role = get_val('director_role', default=sig.get('directorRole', 'PROGRAM DIRECTOR, BSIT')).upper()
    dean_name = get_val('dean_name', default=sig.get('deanName', '')).upper()
    dean_role = get_val('dean_role', default=sig.get('deanRole', 'DEAN, COLLEGE OF COMPUTER STUDIES')).upper()

    # Clean redundant empty paragraphs at the top of the body before transmittal date
    while len(doc.paragraphs) > 0 and not doc.paragraphs[0].text.strip():
        p_elm = doc.paragraphs[0]._element
        p_elm.getparent().remove(p_elm)

    receipts_mode = data.get('receipts_mode') or data.get('receiptsMode') or 'NONE'
    receipts = data.get('receipts') or []

    if receipts_mode == 'STANDALONE':
        for tbl in list(doc.tables):
            t_elm = tbl._element
            t_elm.getparent().remove(t_elm)
        for p in list(doc.paragraphs):
            p_elm = p._element
            p_elm.getparent().remove(p_elm)

        add_receipt_sheets_table(doc, receipts, academic_year, semester, font_name=font_name, font_size_pt=font_size_pt)
        setup_footer(doc.sections[0], style=page_number_style, font_name=font_name)
        doc.save(output_path)
        print(f"Successfully generated standalone receipts report: {output_path}")
        return output_path

    # Ledger Inflows & Outflows
    txs = data.get('transactions', []) or []
    income_items = data.get('income_items')
    if income_items is None:
        income_items = [
            {'title': t.get('title') or 'Income', 'amount': float(t.get('amount') or 0)}
            for t in txs if t.get('type') == 'INFLOW' and 'rollover' not in (t.get('category_name') or '').lower() and 'beginning' not in (t.get('title') or '').lower()
        ]

    expense_items = data.get('expense_items')
    if expense_items is None:
        expense_items = [
            {'title': t.get('title') or 'Expense', 'amount': float(t.get('amount') or 0)}
            for t in txs if t.get('type') == 'OUTFLOW'
        ]

    initial_budget = float(data.get('initial_budget', summary.get('initial_budget', 0.0)) or 0.0)
    total_income = sum(float(i.get('amount', 0)) for i in income_items)
    total_funds = initial_budget + total_income
    total_expenses = sum(float(i.get('amount', 0)) for i in expense_items)
    cash_on_hand = total_funds - total_expenses

    config = {
        '{{TRANSMITTAL_DATE}}': transmittal_date,
        '{{COORDINATOR_NAME}}': coordinator_name,
        '{{COORDINATOR_ROLE}}': coordinator_role,
        '{{COORDINATOR_SALUTATION}}': coordinator_salutation,
        '{{SEMESTER}}': semester,
        '{{ACADEMIC_YEAR}}': academic_year,
        '{{PERIOD_DESCRIPTION}}': period_desc,
        '{{FINAL_AS_OF_DATE}}': final_as_of,
        
        # Totals
        '{{INITIAL_BUDGET_AMOUNT}}': f"P{initial_budget:,.2f}",
        '{{TOTAL_INCOME_AMOUNT}}': f"P{total_income:,.2f}",
        '{{TOTAL_EXPENSES_AMOUNT}}': f"P{total_expenses:,.2f}",
        '{{TOTAL_FUNDS_AMOUNT}}': f"P{total_funds:,.2f}",
        '{{CASH_ON_HAND_AMOUNT}}': f"P{cash_on_hand:,.2f}"
    }

    # Replace placeholders in paragraphs and tables
    replace_placeholders(doc, config)

    # -------------------------------------------------------------
    # 2. TABLE 0: INITIAL BUDGET (Carried over from prior semester)
    # -------------------------------------------------------------
    tbl0 = doc.tables[0]
    ensure_invisible_table_borders(tbl0)
    tbl0.rows[0].cells[0].text = 'INITIAL BUDGET / BEGINNING BALANCE (Carried over from prior semester)'
    for r in tbl0.rows[0].cells[0].paragraphs[0].runs:
        r.font.name = font_name
        r.font.size = Pt(font_size_pt)
        r.bold = True
    tbl0.rows[1].cells[4].text = f"P{initial_budget:,.2f}"

    # -------------------------------------------------------------
    # 3. TABLE 1: INCOME SCHEDULE (STRICTLY FROM DATABASE)
    # -------------------------------------------------------------
    tbl1 = doc.tables[1]
    ensure_invisible_table_borders(tbl1)
    row_template = copy.deepcopy(tbl1.rows[1]._element)
    
    # Remove all template dummy rows
    while len(tbl1.rows) > 1:
        tr = tbl1.rows[-1]._element
        tr.getparent().remove(tr)

    if not income_items:
        new_tr = copy.deepcopy(row_template)
        tbl1._element.append(new_tr)
        row = tbl1.rows[-1]
        row.cells[0].text = "No income transactions recorded."
        row.cells[1].text = ""
        row.cells[2].text = "P0.00"
    else:
        for item in income_items:
            new_tr = copy.deepcopy(row_template)
            tbl1._element.append(new_tr)
            row = tbl1.rows[-1]
            title = str(item.get('title') or item.get('particulars') or item.get('desc') or 'Income')
            amt = float(item.get('amount', 0) or 0)
            row.cells[0].text = title
            row.cells[1].text = ""
            row.cells[2].text = f"P{amt:,.2f}"

    # Append Total Income summary row
    new_tr = copy.deepcopy(row_template)
    tbl1._element.append(new_tr)
    total_row = tbl1.rows[-1]
    total_row.cells[0].text = "Total Income:"
    total_row.cells[1].text = ""
    total_row.cells[2].text = f"P{total_income:,.2f}"
    set_cell_border(total_row.cells[2], bottom={"val": "single", "sz": "4", "color": "000000"})
    for p in total_row.cells[0].paragraphs:
        for r in p.runs:
            r.bold = True
    for p in total_row.cells[2].paragraphs:
        for r in p.runs:
            r.bold = True

    # -------------------------------------------------------------
    # 4. REMOVE TABLE 2 (DUMMY GAME EXHIBIT SUB-TABLE FROM OLD TEMPLATE)
    # -------------------------------------------------------------
    if len(doc.tables) > 4:
        tbl2 = doc.tables[2]
        tbl2._element.getparent().remove(tbl2._element)

    # -------------------------------------------------------------
    # 5. TABLE 3 (NOW TABLE 2): MISCELLANEOUS EXPENSES
    # -------------------------------------------------------------
    tbl_exp = doc.tables[2]
    ensure_invisible_table_borders(tbl_exp)
    exp_template = copy.deepcopy(tbl_exp.rows[2]._element)

    while len(tbl_exp.rows) > 2:
        tr = tbl_exp.rows[-1]._element
        tr.getparent().remove(tr)

    if not expense_items:
        new_tr = copy.deepcopy(exp_template)
        tbl_exp._element.append(new_tr)
        row = tbl_exp.rows[-1]
        row.cells[0].text = "No expense transactions recorded."
        row.cells[1].text = ""
        row.cells[2].text = "P0.00"
    else:
        for item in expense_items:
            new_tr = copy.deepcopy(exp_template)
            tbl_exp._element.append(new_tr)
            row = tbl_exp.rows[-1]
            title = str(item.get('title') or item.get('particulars') or item.get('desc') or 'Expense')
            amt = float(item.get('amount', 0) or 0)
            row.cells[0].text = title
            row.cells[1].text = ""
            row.cells[2].text = f"P{amt:,.2f}"

    # Append Total Expenses summary row
    new_tr = copy.deepcopy(exp_template)
    tbl_exp._element.append(new_tr)
    total_exp_row = tbl_exp.rows[-1]
    total_exp_row.cells[0].text = "Total Expenses:"
    total_exp_row.cells[1].text = ""
    total_exp_row.cells[2].text = f"P{total_expenses:,.2f}"
    set_cell_border(total_exp_row.cells[2], bottom={"val": "single", "sz": "4", "color": "000000"})
    for p in total_exp_row.cells[0].paragraphs:
        for r in p.runs:
            r.bold = True
    for p in total_exp_row.cells[2].paragraphs:
        for r in p.runs:
            r.bold = True

    # -------------------------------------------------------------
    # 6. TABLE 4 (NOW TABLE 3): RECONCILIATION SUMMARY
    # -------------------------------------------------------------
    tbl_summary = doc.tables[3]
    ensure_invisible_table_borders(tbl_summary)
    tbl_summary.rows[1].cells[2].text = f"P{total_funds:,.2f}"
    tbl_summary.rows[2].cells[2].text = f"P{total_expenses:,.2f}"
    tbl_summary.rows[3].cells[2].text = f"P{cash_on_hand:,.2f}"
    set_cell_border(tbl_summary.rows[2].cells[2], bottom={"val": "single", "sz": "4", "color": "000000"})
    set_cell_border(tbl_summary.rows[3].cells[2], top={"val": "single", "sz": "4", "color": "000000"})

    # -------------------------------------------------------------
    # 7. STRUCTURED SIGNATORIES (PAGE 1 & PAGE 4)
    # -------------------------------------------------------------
    inc = data.get('included_signatories') or data.get('includedSignatories') or {}
    officers = []
    if inc.get('treasurer', True) and tres_name:
        officers.append({'name': tres_name, 'role': tres_role})
    if inc.get('auditor', True) and aud_name:
        officers.append({'name': aud_name, 'role': aud_role})
    if inc.get('vice_president', inc.get('vicePresident', bool(vp_name))) and vp_name:
        officers.append({'name': vp_name, 'role': vp_role})
    if inc.get('president', True) and pres_name:
        officers.append({'name': pres_name, 'role': pres_role})

    advisers = []
    if inc.get('adviser1', inc.get('adviser', True)) and adv1_name:
        advisers.append({'name': adv1_name, 'role': adv1_role})
    if inc.get('adviser2', True) and adv2_name and adv2_name != adv1_name:
        advisers.append({'name': adv2_name, 'role': adv2_role})

    admins = []
    if inc.get('director', True) and dir_name:
        admins.append({'name': dir_name, 'role': dir_role})
    if inc.get('dean', True) and dean_name:
        admins.append({'name': dean_name, 'role': dean_role})

    last_p1_note_tbl = None

    # Single "Prepared by:" section for all student officers
    prep_nodes = [p for p in doc.paragraphs if p.text.strip() == 'Prepared by:']
    for idx, p in enumerate(prep_nodes):
        if len(officers) > 0:
            p.text = 'Prepared by:'
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if p.runs:
                p.runs[0].font.name = font_name
                p.runs[0].font.size = Pt(font_size_pt)
                p.runs[0].bold = True
            tbl = create_borderless_signatory_table(doc, officers, min(max(len(officers), 3), 4), font_name=font_name, font_size_pt=font_size_pt)
            if tbl is not None:
                p._element.addnext(tbl)
        else:
            p._element.getparent().remove(p._element)

    # Completely remove "Approved by:" headers anywhere in the report
    appr_nodes = [p for p in doc.paragraphs if p.text.strip() == 'Approved by:']
    for p in appr_nodes:
        try:
            p._element.getparent().remove(p._element)
        except Exception:
            pass

    # Single "Noted by:" section covering Club Advisers down to Program Director and Dean
    noted_signatories = advisers + admins
    note_nodes = [p for p in doc.paragraphs if p.text.strip() == 'Noted by:']
    for idx, p in enumerate(note_nodes):
        if len(noted_signatories) > 0:
            p.text = 'Noted by:'
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if p.runs:
                p.runs[0].font.name = font_name
                p.runs[0].font.size = Pt(font_size_pt)
                p.runs[0].bold = True
            tbl = create_borderless_noted_table(doc, advisers, admins, font_name=font_name, font_size_pt=font_size_pt)
            if tbl is not None:
                p._element.addnext(tbl)
                if idx == 0:
                    last_p1_note_tbl = tbl
        else:
            p._element.getparent().remove(p._element)

    # Remove old template placeholder paragraphs for signatories
    body = doc._element.body
    for p in list(doc.paragraphs):
        if any(tag in p.text for tag in [
            '{{PRESIDENT_NAME}}', '{{PRESIDENT_ROLE}}',
            '{{TREASURER_NAME}}', '{{TREASURER_ROLE}}',
            '{{AUDITOR_NAME}}', '{{AUDITOR_ROLE}}',
            '{{ADVISER_NAME}}', '{{ADVISER_ROLE}}',
            '{{DIRECTOR_NAME}}', '{{DIRECTOR_ROLE}}',
            '{{DEAN_NAME}}', '{{DEAN_ROLE}}'
        ]):
            try:
                body.remove(p._element)
            except Exception:
                pass

    # -------------------------------------------------------------
    # 8. CC ROUTING BLOCK & PAGE BREAK AFTER TRANSMITTAL LETTER (PAGE 1)
    # -------------------------------------------------------------
    cc_paragraphs = [p for p in doc.paragraphs if 'cc:' in p.text.lower()]
    for p in cc_paragraphs:
        try:
            body.remove(p._element)
        except Exception:
            pass

    if not include_cc:
        if last_p1_note_tbl is not None:
            p_break = parse_xml(r'<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls('w'))
            last_p1_note_tbl.addnext(p_break)
    else:
        anchor = last_p1_note_tbl
        half_pt_val = int(font_size_pt * 2)
        for rec in cc_recipients:
            if rec.strip():
                p_cc = parse_xml(r'<w:p %s><w:pPr><w:spacing w:before="60" w:after="20"/><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="%d"/><w:color w:val="000000"/></w:rPr></w:pPr><w:r><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="%d"/><w:color w:val="000000"/></w:rPr><w:t>%s</w:t></w:r></w:p>' % (nsdecls('w'), font_name, font_name, half_pt_val, font_name, font_name, half_pt_val, rec.strip()))
                anchor.addnext(p_cc)
                anchor = p_cc
        p_break = parse_xml(r'<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls('w'))
        anchor.addnext(p_break)

    # -------------------------------------------------------------
    # 9. SECTION FLOW & PAGE BREAKS (PAGE 2 -> PAGE 3 -> PAGE 4)
    # -------------------------------------------------------------
    # Remove empty spacer paragraphs between Table 1 and Table 2 (Expenses)
    in_gap = False
    for el in list(body):
        if el is tbl1._element:
            in_gap = True
            continue
        if el is tbl_exp._element:
            in_gap = False
            break
        if in_gap and el.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(el, doc)
            if not p.text.strip():
                body.remove(el)

    tbl_exp_p = parse_xml(r'<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls('w'))
    tbl_exp._element.addprevious(tbl_exp_p)

    # Remove empty spacer paragraphs between Table 2 (Expenses) and Final As Of Date
    final_p = None
    for p in doc.paragraphs:
        if final_as_of in p.text:
            final_p = p
            break

    if final_p is not None:
        in_gap = False
        for el in list(body):
            if el is tbl_exp._element:
                in_gap = True
                continue
            if el is final_p._element:
                in_gap = False
                break
            if in_gap and el.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(el, doc)
                if not p.text.strip():
                    body.remove(el)

        p_break = parse_xml(r'<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % nsdecls('w'))
        final_p._element.addprevious(p_break)

    # Clean consecutive/excessive empty paragraphs across the entire body
    prev_empty = False
    for el in list(body):
        if el.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(el, doc)
            has_br = any('w:br' in r._r.xml for r in p.runs)
            is_empty = not p.text.strip() and not has_br
            if is_empty:
                if prev_empty:
                    body.remove(el)
                else:
                    prev_empty = True
            else:
                prev_empty = False
        else:
            prev_empty = False

    # -------------------------------------------------------------
    # 10. STRICT TYPOGRAPHY ENFORCEMENT & BORDERLESS TABLE STYLING
    # -------------------------------------------------------------
    enforce_calibri_11(doc, font_name=font_name, font_size_pt=font_size_pt)

    # -------------------------------------------------------------
    # 11. ATTACHED RECEIPT SHEETS (PAGE 5+)
    # -------------------------------------------------------------
    if receipts_mode == 'APPEND' and receipts:
        add_receipt_sheets_table(doc, receipts, academic_year, semester, font_name=font_name, font_size_pt=font_size_pt)

    doc.save(output_path)
    print(f"Successfully generated institutional financial report: {output_path}")
    return output_path

if __name__ == '__main__':
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
        output_file = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
        with open(input_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        populate_report(data, output_path=output_file)
    else:
        cur_td, cur_sem, cur_ay = get_current_academic_calendar()
        sample = {
            'transmittal_date': cur_td,
            'semester': cur_sem,
            'academic_year': cur_ay,
            'period_desc': 'Club week',
            'include_cc': False,
            'page_number_style': 'page_x_of_y',
            'income_items': [
                {'title': 'ID Lace lanyard', 'amount': 8900.0}
            ],
            'expense_items': [
                {'title': 'Candy for Officers', 'amount': 10.0}
            ]
        }
        populate_report(sample)
